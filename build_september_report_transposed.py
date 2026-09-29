#!/usr/bin/env python3
"""
Transposed variant of the September report: BRANDS as rows, DAYS as columns.

Same premium look as the day-wise report (black + red accent, Indian numbering,
right-aligned figures). Reuses extraction + styling helpers from
build_september_report.py.

Output:
- One sheet per zone: rows = brands, columns = days (01 Sep .. 27/28 Sep),
  a TOTAL column and SHARE % column on the right, and a DAILY TOTAL row at the
  bottom with the zone grand total.
- A cover 'SUMMARY' sheet comparing all zones (same as the other report).
"""

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from build_september_report import (
    extract, inr, sty, title_band, set_widths, build_summary_sheet,
    ZONES, SRC, INR, PCT,
    INK, INK2, REDF, REDT, GREY, GREYD, INKTX, LINE, BAND, FILL2, WHITE, FONT, box,
)

OUT = "SEPTEMBER 2026 - ZONE SALES REPORT (BRANDS x DAYS).xlsx"


def build_zone_sheet_T(wb, zone, brands, rows):
    ws = wb.create_sheet(zone.strip())
    ws.sheet_view.showGridLines = False

    ndays = len(rows)
    brand_col = 1
    day_c0 = 2                    # first day column
    total_col = day_c0 + ndays    # brand grand total
    share_col = total_col + 1     # share %
    ncol = share_col

    subtitle = "Brand-wise x Day-wise Sales   \u2022   Amounts in \u20b9 (INR)   \u2022   Recomputed from daily data"
    title_band(ws, ncol, zone, subtitle)

    # ---- header row (days across the top) ----
    hr = 5
    sty(ws.cell(hr, brand_col), bold=True, size=10.5, color=WHITE, fill=INK2,
        align="left", wrap=True, indent=1).value = "BRAND"
    for i, row in enumerate(rows):
        wknd = row["day"] in ("Saturday", "Sunday")
        c = ws.cell(hr, day_c0 + i)
        sty(c, bold=True, size=9, color=(REDT if wknd else WHITE), fill=INK2,
            align="center", wrap=True).value = f"{row['date'].strftime('%d %b')}\n{row['day'][:3]}"
    sty(ws.cell(hr, total_col), bold=True, size=10, color=WHITE, fill=REDF,
        align="center", wrap=True).value = "TOTAL"
    sty(ws.cell(hr, share_col), bold=True, size=9.5, color=WHITE, fill=INK2,
        align="center", wrap=True).value = "SHARE %"
    ws.row_dimensions[hr].height = 36

    # ---- per-day totals + per-brand totals ----
    day_tot = [sum(r["values"].get(b, 0.0) for b in brands) for r in rows]
    brand_tot = {b: sum(r["values"].get(b, 0.0) for r in rows) for b in brands}
    grand = sum(brand_tot.values())
    top_b = max(brand_tot, key=brand_tot.get)

    # ---- data rows (one per brand) ----
    r = hr + 1
    for bi, b in enumerate(brands):
        rf = BAND if bi % 2 else WHITE
        sty(ws.cell(r, brand_col), bold=True, size=10.5, color=INKTX, fill=rf,
            align="left", indent=1).value = b.strip()
        for i, row in enumerate(rows):
            v = row["values"].get(b, 0.0)
            cc = ws.cell(r, day_c0 + i)
            if v:
                sty(cc, size=10, color=INKTX, fill=rf, align="right", fmt=INR).value = v
            else:
                sty(cc, size=10, color="C4C4C4", fill=rf, align="center").value = "\u2013"
        sty(ws.cell(r, total_col), bold=True, size=10.5, color=INKTX, fill=FILL2,
            align="right", fmt=INR, indent=1).value = brand_tot[b]
        share = brand_tot[b] / grand if grand else 0
        is_top = (b == top_b)
        sty(ws.cell(r, share_col), bold=is_top, size=10,
            color=(REDF if is_top else GREYD), fill=FILL2, align="right",
            fmt=PCT, indent=1).value = share
        ws.row_dimensions[r].height = 20
        r += 1

    # ---- DAILY TOTAL row (bottom) ----
    sty(ws.cell(r, brand_col), bold=True, size=10.5, color=WHITE, fill=INK,
        align="left", indent=1).value = "DAILY TOTAL"
    for i, dt_ in enumerate(day_tot):
        sty(ws.cell(r, day_c0 + i), bold=True, size=9.5, color=WHITE, fill=INK,
            align="right", fmt=INR).value = dt_
    sty(ws.cell(r, total_col), bold=True, size=12, color=WHITE, fill=REDF,
        align="right", fmt=INR, indent=1).value = grand
    sty(ws.cell(r, share_col), bold=True, size=10, color=WHITE, fill=INK,
        align="right", fmt=PCT, indent=1).value = 1 if grand else 0
    ws.row_dimensions[r].height = 26

    # ---- widths ----
    widths = {brand_col: max(max((len(b.strip()) for b in brands), default=8) + 3, 20)}
    # widest per-day figure (a daily total) governs the day-column width
    day_w = max(len(inr(max(day_tot))) + 2, 12) if ndays else 12
    for i in range(ndays):
        widths[day_c0 + i] = day_w
    widths[total_col] = max(len(inr(grand)) + 5, 15)
    widths[share_col] = 10
    set_widths(ws, widths)

    ws.freeze_panes = ws.cell(hr + 1, day_c0)   # keep brand column + header visible

    # ---- summary block ----
    last = get_column_letter(min(ncol, 8))
    sr = r + 2
    ws.merge_cells(start_row=sr, start_column=1, end_row=sr, end_column=min(ncol, 8))
    sty(ws.cell(sr, 1), bold=True, size=12.5, color=WHITE, fill=INK, align="left",
        border=None, indent=1).value = "ZONE SUMMARY  \u2014  September 2026"
    ws.row_dimensions[sr].height = 26
    ws.merge_cells(start_row=sr + 1, start_column=1, end_row=sr + 1, end_column=min(ncol, 8))
    ws.cell(sr + 1, 1).fill = PatternFill("solid", fgColor=REDF)
    ws.row_dimensions[sr + 1].height = 3

    best_i = max(range(ndays), key=lambda i: day_tot[i]) if ndays else 0
    metrics = [
        ("Zone Grand Total", inr(grand)),
        ("Operating Days", f"{ndays} days"),
        ("Average Daily Sales", inr(grand / ndays) if ndays else "\u20b90"),
        ("Best Sales Day", f"{rows[best_i]['date'].strftime('%d %b')} ({rows[best_i]['day']})  \u2022  {inr(day_tot[best_i])}"),
        ("Top Brand", f"{top_b.strip()}  \u2022  {inr(brand_tot[top_b])}  ({brand_tot[top_b]/grand*100:.1f}%)"),
        ("Active Brands", str(len(brands))),
    ]
    mr = sr + 2
    for i, (label, val) in enumerate(metrics):
        sty(ws.cell(mr, 1), bold=True, size=10.5, color=GREYD, fill=FILL2,
            align="left", indent=1).value = label
        ws.merge_cells(start_row=mr, start_column=2, end_row=mr, end_column=min(ncol, 8))
        sty(ws.cell(mr, 2), bold=(i == 0), size=(12 if i == 0 else 10.5),
            color=(REDF if i == 0 else INKTX), align="left", indent=1).value = val
        for cc in range(3, min(ncol, 8) + 1):
            ws.cell(mr, cc).border = box
        ws.row_dimensions[mr].height = 21
        mr += 1

    mr += 1
    ws.merge_cells(start_row=mr, start_column=1, end_row=mr, end_column=min(ncol, 8))
    sty(ws.cell(mr, 1), size=9, italic=True, color=GREY, align="left",
        border=None, indent=1).value = (
        "Weekends shown in red \u2022 \u2013 denotes no recorded sale \u2022 "
        "all totals recomputed from daily brand-wise figures")

    return {"grand": grand, "brand_tot": brand_tot, "ndays": ndays,
            "top_brand": top_b.strip(), "best": (rows[best_i]['date'], day_tot[best_i]),
            "brands": [b.strip() for b in brands]}


def main():
    src = load_workbook(SRC, data_only=True)
    out = Workbook()
    out.remove(out.active)
    results = {}
    for zone, meta in ZONES.items():
        brands, rows = extract(src[zone], meta["start"])
        print(f"[{zone.strip()}] brands={len(brands)} days={len(rows)}")
        results[zone] = build_zone_sheet_T(out, zone, brands, rows)
    build_summary_sheet(out, results)
    out.save(OUT)
    print("Saved:", OUT)
    tot = 0
    for zone, d in results.items():
        tot += d["grand"]
        print(f"  {zone.strip():12s} grand={inr(d['grand'])}")
    print("  ALL ZONES   ", inr(tot))


if __name__ == "__main__":
    main()
