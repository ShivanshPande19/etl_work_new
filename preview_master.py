#!/usr/bin/env python3
"""Approximate PNG preview of the beautified master (one month block)."""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
import openpyxl, datetime as dt

INK="#34515B"; INK2="#4A6D76"; INKTX="#2F3E45"; GREYD="#5C6B72"
LINE="#D9D3C8"; BAND="#F7F5F0"; TOTF="#DCEAE4"; WKND="#F1EDE4"
SAGE="#4F7D6E"; WHITE="#FFFFFF"; DASH="#C4C0B6"
MONTHS={m.upper():i for i,m in enumerate(["january","february","march","april","may","june","july","august","september","october","november","december"],1)}
MONTHS.update({m[:3].upper():i for m,i in list(MONTHS.items())})

def inr(v):
    if not isinstance(v,(int,float)) or v==0: return ""
    n=int(round(v)); s=str(n)
    if len(s)>3:
        last3=s[-3:]; rest=s[:-3]; parts=[]
        while len(rest)>2: parts.insert(0,rest[-2:]); rest=rest[:-2]
        if rest: parts.insert(0,rest)
        s=",".join(parts)+","+last3
    return "\u20b9"+s

def render(sheet, month, path, title):
    wb=openpyxl.load_workbook("MASTER SALES REPORT -2.xlsx", data_only=True)
    ws=wb[sheet]; ncol=ws.max_column
    total_col=max((c for c in range(1,ncol+1) if ws.cell(1,c).value and "TOTAL" in str(ws.cell(1,c).value).upper()), default=ncol)
    header=["MONTH","DATE","DAY"]+[str(ws.cell(1,c).value or "").strip().upper() for c in range(4,total_col+1)]
    # find block
    mrows=[(r,MONTHS.get(str(ws.cell(r,1).value).strip().upper())) for r in range(2,ws.max_row+1) if MONTHS.get(str(ws.cell(r,1).value).strip().upper())]
    sr=er=mon=None
    for i,(r,mn) in enumerate(mrows):
        if str(ws.cell(r,1).value).strip().upper().startswith(month.upper()):
            sr=r; mon=mn; er=(mrows[i+1][0]-1) if i+1<len(mrows) else ws.max_row; break
    while er>sr and all(ws.cell(er,c).value in (None,"") for c in range(1,total_col+1)): er-=1
    year=next((ws.cell(r,2).value.year for r in range(sr,er+1) if isinstance(ws.cell(r,2).value,(dt.datetime,dt.date))),2026)

    rows=[]; kinds=[]; di=0
    for r in range(sr,er+1):
        day=ws.cell(r,3).value
        daily=day not in (None,"")
        has=any(ws.cell(r,c).value not in (None,"") for c in range(4,total_col+1))
        if not daily and not has: continue
        if daily:
            di+=1
            try: datestr=dt.date(year,mon,di).strftime("%d %b")
            except: datestr=""
        else:
            datestr=""
        vals=[inr(ws.cell(r,c).value) for c in range(4,total_col+1)]
        rows.append(["", datestr, ("" if not daily else str(day).title()), *vals])
        kinds.append("total" if (not daily and has) else ("wknd" if str(day).strip().upper() in ("SATURDAY","SUNDAY") else "day"))

    nrows=len(rows)+1
    fig_w=min(1.35*len(header)+1, 22); fig_h=0.36*nrows+1.4
    fig,ax=plt.subplots(figsize=(fig_w,fig_h)); ax.axis("off")
    ax.set_title(title, fontsize=14, fontweight="bold", color=INKTX, pad=14, loc="left")
    tbl=ax.table(cellText=rows, colLabels=header, cellLoc="center", loc="center", bbox=[0,0,1,0.95])
    tbl.auto_set_font_size(False); tbl.set_fontsize(8.5); tbl.scale(1,1.25)
    ncols=len(header)
    for (rr,cc),cell in tbl.get_celld().items():
        cell.set_edgecolor(LINE); cell.set_linewidth(0.6); txt=cell.get_text()
        if rr==0:
            cell.set_facecolor(INK2); txt.set_color("white"); txt.set_fontweight("bold"); txt.set_fontsize(8)
            cell.set_height(cell.get_height()*1.4); continue
        k=kinds[rr-1]
        if cc==0:  # MONTH band
            cell.set_facecolor(INK)
            continue
        if k=="total": cell.set_facecolor(TOTF)
        elif k=="wknd": cell.set_facecolor(WKND)
        else: cell.set_facecolor(BAND if (rr-1)%2 else WHITE)
        if cc==1: txt.set_color(INKTX)
        elif cc==2: txt.set_color(SAGE if k=="wknd" else GREYD); txt.set_fontweight("bold" if k=="wknd" else "normal")
        else:
            txt.set_ha("right")
            txt.set_color(INK if k=="total" else INKTX)
        if k=="total": txt.set_fontweight("bold")
    # month name rotated over the band column
    ax.text(0.5/ncols, 0.475, header and "", transform=ax.transAxes)
    ax.text(0.5/ncols, 0.47, f"{sheet.strip()}  \u00b7  {month.upper()} {year}", transform=ax.transAxes,
            rotation=90, ha="center", va="center", color="white", fontsize=11, fontweight="bold")
    plt.savefig(path,dpi=150,bbox_inches="tight",facecolor="white"); print("wrote",path)

render("CENTRAL 50","DECEMBER","preview_master_central.png","Master file — formatted (CENTRAL 50)")
render("BENNETT ","AUG","preview_master_bennett.png","Master file — formatted (BENNETT, 9 brands)")
