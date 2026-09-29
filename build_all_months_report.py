#!/usr/bin/env python3
"""
ALL-MONTHS zone-wise DAILY BRAND SALES report, PetPooja colour scheme.

One sheet per zone. Inside each sheet, every month's daily-brand table is
stacked one below the other (ek ke neeche ek), each with its own header band
and month TOTAL row, then a final ZONE GRAND TOTAL across all months.

Layout per month block:
  DATE | DAY | <brand columns...> | DAILY TOTAL
Design: Office-blue band (#2E5A87, white text), peach weekend rows
(#FCE4D6, day in red #C00000), light-yellow TOTAL rows (#FFF2CC),
light-blue daily-total column (#D9E1F2), Calibri, Indian rupee format,
en-dash for zero, frozen header columns, gridlines hidden.
All totals recomputed from daily brand-wise figures.
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from extract_all_months import load_all

OUT = "MONTHLY SALES REPORT - ALL ZONES (PetPooja style).xlsx"
FONT = "Calibri"

# ---- PetPooja palette ----
BLUE   = "2E5A87"   # header band / grand-total fill (white text)
LBLUE  = "D9E1F2"   # daily-total column fill
PEACH  = "FCE4D6"   # weekend row fill
YELLOW = "FFF2CC"   # total row fill
RED    = "C00000"   # accent text (grand total / weekend day)
GREY   = "666666"   # secondary text
INKTX  = "1F1F1F"   # primary text
LINE   = "BFBFBF"   # thin borders
DASHC  = "A6A6A6"   # dash for zero

INR = '[>=10000000]"\u20b9"##\\,##\\,##\\,##0;[>=100000]"\u20b9"##\\,##\\,##0;"\u20b9"#,##0'

hair = Side(style="thin", color=LINE)
box = Border(left=hair, right=hair, top=hair, bottom=hair)
med = Side(style="medium", color=BLUE)


def inr(n):
    n = int(round(n)); sign = "-" if n < 0 else ""; s = str(abs(n))
    if len(s) > 3:
        last3 = s[-3:]; rest = s[:-3]; parts = []
        while len(rest) > 2:
            parts.insert(0, rest[-2:]); rest = rest[:-2]
        if rest:
            parts.insert(0, rest)
        grouped = ",".join(parts) + "," + last3
    else:
        grouped = s
    return f"{sign}\u20b9{grouped}"


def sty(cl, *, bold=False, size=11, color=INKTX, fill=None, align="center",
        fmt=None, border=box, italic=False, wrap=False, indent=0):
    cl.font = Font(name=FONT, bold=bold, size=size, color=color, italic=italic)
    cl.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap, indent=indent)
    if fill:
        cl.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        cl.number_format = fmt
    if border is not None:
        cl.border = border
    return cl


def month_block(ws, r, month, brands, ncol, last):
    """Render one month's daily-brand table starting at row r. Returns (next_row, month_total, brand_totals)."""
    dtot = ncol
    name = month["name"].title()
    yr = month["year"]

    # month sub-title band (blue fill, white text)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    sty(ws.cell(r, 1), bold=True, size=13, color="FFFFFF", fill=BLUE, align="left",
        indent=1).value = f"{name} {yr}"
    ws.row_dimensions[r].height = 22
    r += 1

    # header
    sty(ws.cell(r, 1), bold=True, size=10.5, color="FFFFFF", fill=BLUE, wrap=True).value = "DATE"
    sty(ws.cell(r, 2), bold=True, size=10.5, color="FFFFFF", fill=BLUE, wrap=True).value = "DAY"
    for j, b in enumerate(brands):
        sty(ws.cell(r, 3 + j), bold=True, size=9.5, color="FFFFFF", fill=BLUE, wrap=True).value = b.strip().upper()
    sty(ws.cell(r, dtot), bold=True, size=10.5, color="FFFFFF", fill=BLUE, wrap=True).value = "DAILY TOTAL"
    ws.row_dimensions[r].height = 30
    r += 1

    brand_tot = {b: 0.0 for b in brands}
    grand = 0.0
    for row in month["rows"]:
        wknd = row["day"] in ("Saturday", "Sunday")
        rf = PEACH if wknd else "FFFFFF"
        sty(ws.cell(r, 1), size=10.5, color=INKTX, fill=rf).value = row["date"].strftime("%d %b")
        sty(ws.cell(r, 2), size=10.5, bold=wknd, color=(RED if wknd else GREY), fill=rf).value = row["day"]
        dsum = 0.0
        for j, b in enumerate(brands):
            v = row["values"].get(b, 0.0)
            dsum += v
            brand_tot[b] += v
            c = ws.cell(r, 3 + j)
            if v:
                sty(c, size=10.5, color=INKTX, fill=rf, align="right", fmt=INR, indent=1).value = v
            else:
                sty(c, size=10.5, color=DASHC, fill=rf, align="center").value = "\u2013"
        sty(ws.cell(r, dtot), bold=True, size=10.5, color=INKTX, fill=LBLUE,
            align="right", fmt=INR, indent=1).value = dsum
        grand += dsum
        ws.row_dimensions[r].height = 17
        r += 1

    # month TOTAL row
    sty(ws.cell(r, 1), bold=True, size=11, color=BLUE, fill=YELLOW,
        border=Border(left=hair, right=hair, top=med, bottom=hair)).value = "TOTAL"
    sty(ws.cell(r, 2), fill=YELLOW,
        border=Border(left=hair, right=hair, top=med, bottom=hair)).value = ""
    for j, b in enumerate(brands):
        sty(ws.cell(r, 3 + j), bold=True, size=10.5, color=INKTX, fill=YELLOW, align="right",
            fmt=INR, indent=1, border=Border(left=hair, right=hair, top=med, bottom=hair)).value = brand_tot[b]
    sty(ws.cell(r, dtot), bold=True, size=12, color="FFFFFF", fill=BLUE, align="right",
        fmt=INR, indent=1, border=Border(left=hair, right=hair, top=med, bottom=med)).value = grand
    ws.row_dimensions[r].height = 22
    r += 1

    # per-month ZONE TOTAL callout (right after this month's table)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(2, ncol - 2))
    sty(ws.cell(r, 1), bold=True, size=11.5, color=BLUE, fill=YELLOW, align="left",
        indent=1).value = f"ZONE TOTAL  \u2014  MONTH OF {month['name'].upper()} {yr}"
    ws.merge_cells(start_row=r, start_column=ncol - 1, end_row=r, end_column=ncol)
    sty(ws.cell(r, ncol - 1), bold=True, size=12.5, color=RED, fill=YELLOW,
        align="right", fmt=INR, indent=1).value = grand
    ws.row_dimensions[r].height = 22
    r += 1
    return r, grand, brand_tot


def build_zone(wb, zone, brands, months):
    ws = wb.create_sheet(zone.strip())
    ws.sheet_view.showGridLines = False
    ncol = 2 + len(brands) + 1
    last = get_column_letter(ncol)

    # zone title
    ws.merge_cells(f"A1:{last}1")
    sty(ws.cell(1, 1), bold=True, size=16, color=BLUE, align="left", border=None,
        indent=1).value = f"{zone.strip()}  \u2014  MONTHLY DAILY-BRAND SALES"
    ws.row_dimensions[1].height = 26
    ws.merge_cells(f"A2:{last}2")
    span = f"{months[0]['name'].title()} {months[0]['year']} \u2013 {months[-1]['name'].title()} {months[-1]['year']}"
    sty(ws.cell(2, 1), size=10, color=GREY, italic=True, align="left", border=None,
        indent=1).value = (f"Day-wise net sales by brand, month by month  \u2022  {span}  \u2022  "
                           "Amounts in \u20b9 (INR)  \u2022  weekend rows shaded")
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 6

    r = 4
    zone_grand = 0.0
    zone_brand = {b: 0.0 for b in brands}
    for m in months:
        r, g, bt = month_block(ws, r, m, brands, ncol, last)
        zone_grand += g
        for b in brands:
            zone_brand[b] += bt[b]
        ws.row_dimensions[r].height = 8  # gap row
        r += 1

    # ZONE GRAND TOTAL band (all months)
    sty(ws.cell(r, 1), bold=True, size=11.5, color="FFFFFF", fill=BLUE).value = "ALL"
    sty(ws.cell(r, 2), bold=True, size=10.5, color="FFFFFF", fill=BLUE, align="left",
        indent=0).value = "MONTHS"
    for j, b in enumerate(brands):
        sty(ws.cell(r, 3 + j), bold=True, size=10.5, color="FFFFFF", fill=BLUE, align="right",
            fmt=INR, indent=1).value = zone_brand[b]
    sty(ws.cell(r, ncol), bold=True, size=12.5, color="FFF2CC", fill=BLUE, align="right",
        fmt=INR, indent=1).value = zone_grand
    ws.row_dimensions[r].height = 26
    r += 1

    # explicit callout
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(2, ncol - 2))
    sty(ws.cell(r, 1), bold=True, size=12, color=BLUE, fill=YELLOW, align="left",
        indent=1).value = f"ZONE GRAND TOTAL  \u2014  {zone.strip()} ({len(months)} months)"
    ws.merge_cells(start_row=r, start_column=ncol - 1, end_row=r, end_column=ncol)
    sty(ws.cell(r, ncol - 1), bold=True, size=13, color=RED, fill=YELLOW,
        align="right", fmt=INR, indent=1).value = zone_grand
    ws.row_dimensions[r].height = 24

    # widths
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 12
    for j, b in enumerate(brands):
        w = max(max((len(x) for x in b.strip().split()), default=6) + 3,
                len(inr(zone_brand[b])) + 3, 13)
        ws.column_dimensions[get_column_letter(3 + j)].width = w
    ws.column_dimensions[last].width = max(len(inr(zone_grand)) + 4, 15)
    ws.freeze_panes = "C5"
    return zone_grand


def main():
    data = load_all()
    wb = Workbook()
    wb.remove(wb.active)
    total = 0.0
    for zone, d in data.items():
        g = build_zone(wb, zone, d["brands"], d["months"])
        total += g
        print(f"  {zone.strip():12s} months={len(d['months']):2d} brands={len(d['brands'])} "
              f"grand={inr(g)}")
    wb.save(OUT)
    print("Saved:", OUT, "| ALL ZONES", inr(total))


if __name__ == "__main__":
    main()
