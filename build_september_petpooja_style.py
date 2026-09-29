#!/usr/bin/env python3
"""
September 2026 — all-zone sales report in the PetPooja report *style*
(Summary / Sales-by-Weekday pivot / Daily Report / Weekday Totals / Zone x Brand
/ Notes), built from the master workbook.

Data reality: the master has only daily NET SALES per brand per zone, so the
order-level PetPooja metrics (bills, avg bill value, payment mode, dine-in vs
pickup, cancellations) are NOT reproduced — only the sales-based views are.
"""

import datetime as dt
from collections import defaultdict
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from build_september_report import (
    extract, inr, sty, title_band, set_widths, SRC, ZONES,
    INK, INK2, REDF, REDT, GREY, GREYD, INKTX, LINE, BAND, FILL2, WHITE, FONT, box, INR, PCT,
)

OUT = "SEPTEMBER 2026 - ALL ZONE REPORT (PetPooja style).xlsx"
YEAR = 2026
WDNAMES = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def ind(v):
    """Indian-grouped number without the rupee sign."""
    return inr(v).lstrip("\u20b9")


def load_zones():
    src = load_workbook(SRC, data_only=True)
    zones = {}
    for zone, meta in ZONES.items():
        brands, rows = extract(src[zone], meta["start"])
        daily = {r["date"]: sum(r["values"].get(b, 0.0) for b in brands) for r in rows}
        btot = {b: sum(r["values"].get(b, 0.0) for r in rows) for b in brands}
        zones[zone.strip()] = {
            "brands": [b.strip() for b in brands],
            "btot": {b.strip(): v for b, v in btot.items()},
            "daily": daily,
            "grand": sum(daily.values()),
            "days": len(daily),
        }
    return zones


# ---------------- Summary ----------------
def sheet_summary(wb, zones):
    ws = wb.create_sheet("Summary")
    ws.sheet_view.showGridLines = False
    ncol = 6
    title_band(ws, ncol, "SUMMARY",
               "All-Zone Sales Overview  \u2022  Net sales in \u20b9  \u2022  Data: master workbook (no order-level metrics)")

    grand_all = sum(z["grand"] for z in zones.values())
    all_daily = defaultdict(float)
    for z in zones.values():
        for d, v in z["daily"].items():
            all_daily[d] += v
    best_day = max(all_daily, key=all_daily.get) if all_daily else None
    top_zone = max(zones, key=lambda z: zones[z]["grand"])
    # top brand overall (brand names can repeat across zones -> aggregate)
    brand_all = defaultdict(float)
    for z in zones.values():
        for b, v in z["btot"].items():
            brand_all[b] += v
    top_brand = max(brand_all, key=brand_all.get)

    r = 5
    kpis = [
        ("All-Zone Net Sales", inr(grand_all)),
        ("Active Zones", f"{len(zones)}  (" + ", ".join(zones) + ")"),
        ("Sales Window", f"01 Sep \u2013 {max(all_daily).strftime('%d %b %Y')}"),
        ("Best Sales Day", f"{best_day.strftime('%d %b')} ({best_day.strftime('%A')})  \u2022  {inr(all_daily[best_day])}"),
        ("Top Zone", f"{top_zone}  \u2022  {inr(zones[top_zone]['grand'])}  ({zones[top_zone]['grand']/grand_all*100:.1f}%)"),
        ("Top Brand (all zones)", f"{top_brand}  \u2022  {inr(brand_all[top_brand])}"),
    ]
    for i, (k, v) in enumerate(kpis):
        sty(ws.cell(r, 1), bold=True, size=10.5, color=GREYD, fill=FILL2, align="left", indent=1).value = k
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=ncol)
        sty(ws.cell(r, 2), bold=(i == 0), size=(12 if i == 0 else 10.5),
            color=(REDF if i == 0 else INKTX), align="left", indent=1).value = v
        for c in range(3, ncol + 1):
            ws.cell(r, c).border = box
        ws.row_dimensions[r].height = 21
        r += 1

    # per-zone table
    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    sty(ws.cell(r, 1), bold=True, size=12, color=WHITE, fill=INK, align="left", border=None, indent=1).value = "BY ZONE"
    ws.row_dimensions[r].height = 24
    r += 1
    heads = ["ZONE", "OPERATING DAYS", "NET SALES", "AVG / DAY", "TOP BRAND", "SHARE %"]
    fills = [INK2, INK2, REDF, INK2, INK2, INK2]
    for c, (h, f) in enumerate(zip(heads, fills), 1):
        sty(ws.cell(r, c), bold=True, size=10, color=WHITE, fill=f, wrap=True).value = h
    ws.row_dimensions[r].height = 30
    r += 1
    for i, (zone, z) in enumerate(sorted(zones.items(), key=lambda kv: -kv[1]["grand"])):
        rf = BAND if i % 2 else WHITE
        tb = max(z["btot"], key=z["btot"].get)
        sty(ws.cell(r, 1), bold=True, size=10.5, color=INKTX, fill=rf, align="left", indent=1).value = zone
        sty(ws.cell(r, 2), size=10.5, color=INKTX, fill=rf).value = f"{z['days']} days"
        sty(ws.cell(r, 3), bold=True, size=10.5, color=REDF, fill=rf, align="right", fmt=INR, indent=1).value = z["grand"]
        sty(ws.cell(r, 4), size=10.5, color=INKTX, fill=rf, align="right", fmt=INR, indent=1).value = z["grand"] / z["days"]
        sty(ws.cell(r, 5), size=10.5, color=INKTX, fill=rf, align="left", indent=1).value = tb
        sty(ws.cell(r, 6), size=10, color=GREYD, fill=rf, align="right", fmt=PCT, indent=1).value = z["grand"] / grand_all
        ws.row_dimensions[r].height = 22
        r += 1
    sty(ws.cell(r, 1), bold=True, size=11, color=WHITE, fill=INK, align="left", indent=1).value = "ALL ZONES"
    for c in range(2, ncol + 1):
        sty(ws.cell(r, c), fill=INK, border=box).value = ""
    sty(ws.cell(r, 3), bold=True, size=12, color=WHITE, fill=REDF, align="right", fmt=INR, indent=1).value = grand_all
    sty(ws.cell(r, 6), bold=True, size=10, color=WHITE, fill=INK, align="right", fmt=PCT, indent=1).value = 1
    ws.row_dimensions[r].height = 26

    set_widths(ws, {1: 16, 2: 16, 3: 16, 4: 14, 5: 22, 6: 10})
    ws.freeze_panes = "A5"


# ---------------- weekday pivot (one table) ----------------
def write_pivot(ws, r0, label, daily, ncol_hint=None):
    dates = sorted(daily)
    weeks = sorted({d - dt.timedelta(days=d.weekday()) for d in dates})
    ncol = 1 + len(weeks) + 2  # weekday + weeks + TOTAL + AVG
    last = get_column_letter(ncol)

    ws.merge_cells(start_row=r0, start_column=1, end_row=r0, end_column=ncol)
    sty(ws.cell(r0, 1), bold=True, size=11.5, color=WHITE, fill=INK, align="left", border=None, indent=1).value = label
    ws.row_dimensions[r0].height = 22
    hr = r0 + 1
    sty(ws.cell(hr, 1), bold=True, size=9.5, color=WHITE, fill=INK2, align="left", indent=1, wrap=True).value = "Weekday"
    for j, wk in enumerate(weeks):
        sty(ws.cell(hr, 2 + j), bold=True, size=9, color=WHITE, fill=INK2, wrap=True).value = f"Week of\n{wk.strftime('%d %b')}"
    sty(ws.cell(hr, ncol - 1), bold=True, size=9.5, color=WHITE, fill=REDF, wrap=True).value = "TOTAL"
    sty(ws.cell(hr, ncol), bold=True, size=9, color=WHITE, fill=INK2, wrap=True).value = "AVG / day"
    ws.row_dimensions[hr].height = 30

    col_tot = defaultdict(float)
    grand = 0.0
    for wi, wd in enumerate(WDNAMES):
        r = hr + 1 + wi
        wknd = wd in ("Saturday", "Sunday")
        sty(ws.cell(r, 1), bold=True, size=10, color=(REDF if wknd else GREYD),
            fill=(FILL2 if wknd else WHITE), align="left", indent=1).value = wd
        rowvals = []
        for j, wk in enumerate(weeks):
            d = wk + dt.timedelta(days=wi)
            v = daily.get(d)
            cell = ws.cell(r, 2 + j)
            if v is not None:
                sty(cell, size=9.5, color=INKTX, fill=(BAND if wi % 2 else WHITE),
                    align="center", wrap=True).value = f"{d.strftime('%d %b')}\n{ind(v)}"
                rowvals.append(v)
                col_tot[j] += v
                grand += v
            else:
                sty(cell, fill=(BAND if wi % 2 else WHITE)).value = None
        tot = sum(rowvals)
        sty(ws.cell(r, ncol - 1), bold=True, size=10, color=INKTX, fill=FILL2,
            align="right", fmt=INR, indent=1).value = tot if rowvals else None
        sty(ws.cell(r, ncol), size=10, color=GREYD, fill=(BAND if wi % 2 else WHITE),
            align="right", fmt=INR, indent=1).value = (tot / len(rowvals)) if rowvals else None
        ws.row_dimensions[r].height = 26

    r = hr + 1 + 7
    sty(ws.cell(r, 1), bold=True, size=10, color=WHITE, fill=INK, align="left", indent=1).value = "TOTAL"
    for j in range(len(weeks)):
        sty(ws.cell(r, 2 + j), bold=True, size=9.5, color=WHITE, fill=INK, align="center", fmt=INR).value = col_tot[j]
    sty(ws.cell(r, ncol - 1), bold=True, size=11, color=WHITE, fill=REDF, align="right", fmt=INR, indent=1).value = grand
    sty(ws.cell(r, ncol), fill=INK, border=box).value = ""
    ws.row_dimensions[r].height = 24
    return r + 2, ncol


def sheet_weekday(wb, zones):
    ws = wb.create_sheet("Sales by Weekday")
    ws.sheet_view.showGridLines = False
    title_band(ws, 8, "WEEKDAY",
               "Net Sales by Weekday  \u2022  each cell = date + that day's sales  \u2022  \u20b9")
    r = 5
    maxcol = 8
    # combined first, then each zone
    all_daily = defaultdict(float)
    for z in zones.values():
        for d, v in z["daily"].items():
            all_daily[d] += v
    r, nc = write_pivot(ws, r, "ALL ZONES COMBINED  \u2014  September 2026", dict(all_daily))
    maxcol = max(maxcol, nc)
    for zone in sorted(zones, key=lambda z: -zones[z]["grand"]):
        r, nc = write_pivot(ws, r, f"{zone}  \u2014  September 2026", zones[zone]["daily"])
        maxcol = max(maxcol, nc)

    widths = {1: 12}
    for c in range(2, maxcol - 1):
        widths[c] = 12
    widths[maxcol - 1] = 14
    widths[maxcol] = 12
    set_widths(ws, widths)
    ws.freeze_panes = "B5"


# ---------------- Daily Report ----------------
def sheet_daily(wb, zones):
    ws = wb.create_sheet("Daily Report")
    ws.sheet_view.showGridLines = False
    znames = sorted(zones, key=lambda z: -zones[z]["grand"])
    ncol = 2 + len(znames) + 1  # Date, Weekday, zones..., Grand Total
    title_band(ws, ncol, "DAILY",
               "Day-wise Net Sales by Zone  \u2022  \u20b9")
    hr = 5
    heads = ["DATE", "WEEKDAY"] + znames + ["GRAND TOTAL"]
    for c, h in enumerate(heads, 1):
        fill = REDF if h == "GRAND TOTAL" else INK2
        sty(ws.cell(hr, c), bold=True, size=9.5, color=WHITE, fill=fill, wrap=True).value = h
    ws.row_dimensions[hr].height = 30

    all_dates = sorted({d for z in zones.values() for d in z["daily"]})
    ztot = defaultdict(float)
    gtot = 0.0
    r = hr + 1
    for i, d in enumerate(all_dates):
        wknd = d.weekday() >= 5
        rf = FILL2 if wknd else (BAND if i % 2 else WHITE)
        sty(ws.cell(r, 1), size=10, color=INKTX, fill=rf, align="center").value = d.strftime("%d %b")
        sty(ws.cell(r, 2), size=10, bold=wknd, color=(REDF if wknd else GREYD), fill=rf, align="center").value = d.strftime("%A")
        row_total = 0.0
        for j, zone in enumerate(znames):
            v = zones[zone]["daily"].get(d)
            cell = ws.cell(r, 3 + j)
            if v is not None:
                sty(cell, size=10, color=INKTX, fill=rf, align="right", fmt=INR, indent=1).value = v
                row_total += v
                ztot[zone] += v
            else:
                sty(cell, size=10, color=GREY, fill=rf, align="center").value = "\u2013"
        sty(ws.cell(r, ncol), bold=True, size=10, color=INK, fill=rf, align="right", fmt=INR, indent=1).value = row_total
        gtot += row_total
        ws.row_dimensions[r].height = 19
        r += 1
    # total row
    sty(ws.cell(r, 1), bold=True, size=10.5, color=WHITE, fill=INK, align="center").value = "TOTAL"
    sty(ws.cell(r, 2), fill=INK, border=box).value = ""
    for j, zone in enumerate(znames):
        sty(ws.cell(r, 3 + j), bold=True, size=10, color=WHITE, fill=INK, align="right", fmt=INR, indent=1).value = ztot[zone]
    sty(ws.cell(r, ncol), bold=True, size=12, color=WHITE, fill=REDF, align="right", fmt=INR, indent=1).value = gtot
    ws.row_dimensions[r].height = 26

    widths = {1: 11, 2: 12}
    for j in range(len(znames)):
        widths[3 + j] = 15
    widths[ncol] = 16
    set_widths(ws, widths)
    ws.freeze_panes = ws.cell(hr + 1, 3)


# ---------------- Weekday Totals ----------------
def sheet_wdtotals(wb, zones):
    ws = wb.create_sheet("Weekday Totals")
    ws.sheet_view.showGridLines = False
    ncol = 4
    title_band(ws, ncol, "WEEKDAY",
               "All-Zone totals per weekday  \u2022  \u20b9")
    all_daily = defaultdict(float)
    for z in zones.values():
        for d, v in z["daily"].items():
            all_daily[d] += v
    by_wd = defaultdict(list)
    for d, v in all_daily.items():
        by_wd[d.weekday()].append(v)

    hr = 5
    heads = ["WEEKDAY", "# OF DAYS", "TOTAL NET SALES", "AVG / DAY"]
    fills = [INK2, INK2, REDF, INK2]
    for c, (h, f) in enumerate(zip(heads, fills), 1):
        sty(ws.cell(hr, c), bold=True, size=10, color=WHITE, fill=f, wrap=True).value = h
    ws.row_dimensions[hr].height = 30
    r = hr + 1
    gt = 0.0
    for wi, wd in enumerate(WDNAMES):
        vals = by_wd.get(wi, [])
        wknd = wd in ("Saturday", "Sunday")
        rf = FILL2 if wknd else (BAND if wi % 2 else WHITE)
        sty(ws.cell(r, 1), bold=True, size=10.5, color=(REDF if wknd else INKTX), fill=rf, align="left", indent=1).value = wd
        sty(ws.cell(r, 2), size=10.5, color=INKTX, fill=rf).value = len(vals)
        sty(ws.cell(r, 3), bold=True, size=10.5, color=REDF, fill=rf, align="right", fmt=INR, indent=1).value = sum(vals)
        sty(ws.cell(r, 4), size=10.5, color=INKTX, fill=rf, align="right", fmt=INR, indent=1).value = (sum(vals)/len(vals)) if vals else None
        gt += sum(vals)
        ws.row_dimensions[r].height = 22
        r += 1
    sty(ws.cell(r, 1), bold=True, size=11, color=WHITE, fill=INK, align="left", indent=1).value = "TOTAL"
    sty(ws.cell(r, 2), fill=INK, border=box).value = ""
    sty(ws.cell(r, 3), bold=True, size=12, color=WHITE, fill=REDF, align="right", fmt=INR, indent=1).value = gt
    sty(ws.cell(r, 4), fill=INK, border=box).value = ""
    ws.row_dimensions[r].height = 26
    set_widths(ws, {1: 16, 2: 12, 3: 18, 4: 14})
    ws.freeze_panes = "A5"


# ---------------- Zone x Brand ----------------
def sheet_zonebrand(wb, zones):
    ws = wb.create_sheet("Zone x Brand")
    ws.sheet_view.showGridLines = False
    ncol = 4
    title_band(ws, ncol, "BRANDS",
               "Brand-wise net sales, grouped by zone  \u2022  \u20b9")
    grand_all = sum(z["grand"] for z in zones.values())
    hr = 5
    heads = ["ZONE", "BRAND", "NET SALES", "SHARE % (all)"]
    fills = [INK2, INK2, REDF, INK2]
    for c, (h, f) in enumerate(zip(heads, fills), 1):
        sty(ws.cell(hr, c), bold=True, size=10, color=WHITE, fill=f, wrap=True).value = h
    ws.row_dimensions[hr].height = 30
    r = hr + 1
    i = 0
    for zone in sorted(zones, key=lambda z: -zones[z]["grand"]):
        z = zones[zone]
        items = sorted(z["btot"].items(), key=lambda kv: -kv[1])
        for k, (b, v) in enumerate(items):
            rf = BAND if i % 2 else WHITE
            sty(ws.cell(r, 1), bold=(k == 0), size=10.5, color=INKTX, fill=rf, align="left", indent=1).value = (zone if k == 0 else "")
            sty(ws.cell(r, 2), size=10.5, color=INKTX, fill=rf, align="left", indent=1).value = b
            sty(ws.cell(r, 3), size=10.5, color=INKTX, fill=rf, align="right", fmt=INR, indent=1).value = v
            sty(ws.cell(r, 4), size=10, color=GREYD, fill=rf, align="right", fmt=PCT, indent=1).value = v / grand_all
            ws.row_dimensions[r].height = 19
            r += 1
            i += 1
        # zone subtotal
        sty(ws.cell(r, 1), fill=FILL2, border=box).value = ""
        sty(ws.cell(r, 2), bold=True, size=10, color=INK, fill=FILL2, align="right", indent=1).value = f"{zone} total"
        sty(ws.cell(r, 3), bold=True, size=10.5, color=INK, fill=FILL2, align="right", fmt=INR, indent=1).value = z["grand"]
        sty(ws.cell(r, 4), bold=True, size=10, color=GREYD, fill=FILL2, align="right", fmt=PCT, indent=1).value = z["grand"]/grand_all
        ws.row_dimensions[r].height = 20
        r += 1
    sty(ws.cell(r, 1), bold=True, size=11, color=WHITE, fill=INK, align="left", indent=1).value = "ALL ZONES"
    sty(ws.cell(r, 2), fill=INK, border=box).value = ""
    sty(ws.cell(r, 3), bold=True, size=12, color=WHITE, fill=REDF, align="right", fmt=INR, indent=1).value = grand_all
    sty(ws.cell(r, 4), bold=True, size=10, color=WHITE, fill=INK, align="right", fmt=PCT, indent=1).value = 1
    ws.row_dimensions[r].height = 26
    set_widths(ws, {1: 16, 2: 24, 3: 16, 4: 14})
    ws.freeze_panes = ws.cell(hr + 1, 1)


# ---------------- Notes ----------------
def sheet_notes(wb, zones):
    ws = wb.create_sheet("Notes & Method")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 110
    lines = [
        ("SEPTEMBER 2026  \u2014  ALL-ZONE SALES REPORT  (PetPooja report style)", "h"),
        ("", ""),
        ("SOURCE", "s"),
        ("Built from 'MASTER SALES REPORT -2.xlsx' \u2014 daily NET SALES per brand per zone.", ""),
        ("Zones with September data: " + ", ".join(zones) + ".", ""),
        ("DISTRICT 9 is excluded \u2014 it has no September data in the master (only Jul/Aug).", ""),
        ("Dates were verified/cleaned to correct calendar dates for September 2026.", ""),
        ("", ""),
        ("WHAT THIS REPORT SHOWS", "s"),
        ("\u2022  Summary       \u2013 all-zone total + per-zone breakdown, top zone/brand, best day.", ""),
        ("\u2022  Sales by Weekday \u2013 weekday x week pivot (cell = date + net sales), combined + per zone.", ""),
        ("\u2022  Daily Report   \u2013 each day's net sales for every zone + grand total.", ""),
        ("\u2022  Weekday Totals \u2013 per-weekday total & average/day (all zones).", ""),
        ("\u2022  Zone x Brand   \u2013 brand-wise net sales grouped by zone, with share %.", ""),
        ("", ""),
        ("NOT INCLUDED (data not available in the master)", "s"),
        ("The PetPooja outlet report is built from order-level data. The master has no such data,", ""),
        ("so these are NOT reproduced: bills/orders count, avg bill value, cancelled orders,", ""),
        ("payment mode (GPay/UPI/Cash/Card), and Dine-In vs Pick-Up split.", ""),
        ("Share these order-level exports for each outlet and the full report can be built.", ""),
        ("", ""),
        ("Amounts are net sales in \u20b9 (Indian numbering). Totals recomputed from daily data.", "i"),
    ]
    r = 1
    for text, kind in lines:
        c = ws.cell(r, 1, text)
        if kind == "h":
            c.font = Font(name=FONT, bold=True, size=15, color=WHITE)
            c.fill = PatternFill("solid", fgColor=INK)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
            ws.row_dimensions[r].height = 30
        elif kind == "s":
            c.font = Font(name=FONT, bold=True, size=11.5, color=REDF)
        elif kind == "i":
            c.font = Font(name=FONT, italic=True, size=10, color=GREY)
        else:
            c.font = Font(name=FONT, size=10.5, color=INKTX)
            c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        r += 1


def main():
    zones = load_zones()
    out = Workbook()
    out.remove(out.active)
    sheet_summary(out, zones)
    sheet_weekday(out, zones)
    sheet_daily(out, zones)
    sheet_wdtotals(out, zones)
    sheet_zonebrand(out, zones)
    sheet_notes(out, zones)
    out.save(OUT)
    print("Saved:", OUT)
    ga = sum(z["grand"] for z in zones.values())
    for z, d in zones.items():
        print(f"  {z:12s} days={d['days']:2d} net={inr(d['grand'])}")
    print("  ALL ZONES  ", inr(ga))


if __name__ == "__main__":
    main()
