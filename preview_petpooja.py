#!/usr/bin/env python3
"""Previews of the PetPooja-style September all-zone report."""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
import openpyxl

WB="SEPTEMBER 2026 - ALL ZONE REPORT (PetPooja style).xlsx"
INK="#34515B"; INK2="#4A6D76"; SAGE="#4F7D6E"; SAGET="#9BC3B6"
GREYD="#5C6B72"; INKTX="#2F3E45"; LINE="#E7E2D9"; BAND="#F7F5F0"
FILL2="#EDF2EF"; WHITE="#FFFFFF"; DASH="#C4C0B6"

def ind(n):
    n=int(round(n)); s=str(abs(n)); sign="-" if n<0 else ""
    if len(s)>3:
        last3=s[-3:]; rest=s[:-3]; parts=[]
        while len(rest)>2: parts.insert(0,rest[-2:]); rest=rest[:-2]
        if rest: parts.insert(0,rest)
        s=",".join(parts)+","+last3
    return sign+s
def rup(n): return "\u20b9"+ind(n)

def cellval(ws,r,c):
    v=ws.cell(r,c).value
    fmt=ws.cell(r,c).number_format or ""
    if isinstance(v,(int,float)):
        if "%" in fmt: return f"{v*100:.1f}%"
        if "\u20b9" in fmt: return rup(v)
        return ind(v)
    return "" if v is None else str(v)

def title_bar(ax, text):
    ax.add_patch(Rectangle((0,1.0),1,0.11,transform=ax.transAxes,color=INK,clip_on=False,zorder=3))
    ax.add_patch(Rectangle((0,0.986),1,0.014,transform=ax.transAxes,color=SAGE,clip_on=False,zorder=3))
    ax.text(0.006,1.055,text[0],transform=ax.transAxes,color=SAGET,fontsize=15,fontweight="bold",va="center",zorder=4)
    ax.text(0.02,1.055,text[1:],transform=ax.transAxes,color="white",fontsize=15,fontweight="bold",va="center",zorder=4)

def render_grid(sheet, r0, r1, ncol, path, title, header_row=None):
    wb=openpyxl.load_workbook(WB); ws=wb[sheet]
    data=[[cellval(ws,r,c) for c in range(1,ncol+1)] for r in range(r0,r1+1)]
    labels0=[data[i][0] for i in range(len(data))]
    nrows=len(data)
    fig_w=min(1.5*ncol+2,20); fig_h=0.42*nrows+1.6
    fig,ax=plt.subplots(figsize=(fig_w,fig_h)); ax.axis("off"); title_bar(ax,title)
    tbl=ax.table(cellText=data, cellLoc="center", loc="center", bbox=[0,0,1,0.92])
    tbl.auto_set_font_size(False); tbl.set_fontsize(9); tbl.scale(1,1.35)
    for (rr,cc),cell in tbl.get_celld().items():
        cell.set_edgecolor(LINE); cell.set_linewidth(0.6); txt=cell.get_text()
        lab=labels0[rr].strip().upper()
        val=data[rr][cc]
        if header_row is not None and (r0+rr)==header_row:
            cell.set_facecolor(INK2); txt.set_color("white"); txt.set_fontweight("bold"); txt.set_fontsize(8.5); continue
        if lab in ("ALL ZONES","TOTAL"):
            cell.set_facecolor(INK); txt.set_color("white"); txt.set_fontweight("bold")
        elif lab.endswith("TOTAL") and lab!="GRAND TOTAL":
            cell.set_facecolor(FILL2); txt.set_fontweight("bold"); txt.set_color(INK)
        else:
            cell.set_facecolor(BAND if rr%2 else WHITE)
            if val=="\u2013": txt.set_color(DASH)
            else: txt.set_color(INKTX)
        if cc==0: txt.set_ha("left")
        elif any(ch.isdigit() for ch in val) and "%" not in val and cc>=2: txt.set_ha("right")
    plt.savefig(path,dpi=150,bbox_inches="tight",facecolor="white"); print("wrote",path)

def render_pivot(path):
    wb=openpyxl.load_workbook(WB); ws=wb["Sales by Weekday"]
    # combined pivot: title row 5, header 6, weekdays 7-13, total 14
    ncol=8
    header=[cellval(ws,6,c) for c in range(1,ncol+1)]
    rows=[[cellval(ws,r,c) for c in range(1,ncol+1)] for r in range(7,15)]
    fig,ax=plt.subplots(figsize=(12,5.5)); ax.axis("off"); title_bar(ax,"Sales by Weekday \u2014 ALL ZONES COMBINED (Sept 2026)")
    tbl=ax.table(cellText=rows, colLabels=header, cellLoc="center", loc="center", bbox=[0,0,1,0.9])
    tbl.auto_set_font_size(False); tbl.set_fontsize(8.5); tbl.scale(1,2.0)
    for (rr,cc),cell in tbl.get_celld().items():
        cell.set_edgecolor(LINE); cell.set_linewidth(0.6); txt=cell.get_text()
        if rr==0:
            cell.set_facecolor(SAGE if cc==ncol-2 else INK2); txt.set_color("white"); txt.set_fontweight("bold"); txt.set_fontsize(8)
            continue
        lab=rows[rr-1][0]
        if lab=="TOTAL":
            cell.set_facecolor(SAGE if cc==ncol-2 else INK); txt.set_color("white"); txt.set_fontweight("bold")
        else:
            wknd=lab in ("Saturday","Sunday")
            cell.set_facecolor(FILL2 if (cc in (0,ncol-2)) else (BAND if rr%2 else WHITE))
            if cc==0: txt.set_color(SAGE if wknd else GREYD); txt.set_fontweight("bold"); txt.set_ha("left")
            elif cc>=ncol-2: txt.set_ha("right"); txt.set_color(INKTX); txt.set_fontweight("bold" if cc==ncol-2 else "normal")
            else: txt.set_color(INKTX)
    plt.savefig(path,dpi=150,bbox_inches="tight",facecolor="white"); print("wrote",path)

# Summary: KPI rows 5-10, by-zone header 13, rows 14-16, all-zones 17
render_grid("Summary",5,17,6,"preview_pp_summary.png","Summary \u2014 All-Zone September 2026",header_row=13)
render_pivot("preview_pp_weekday.png")
# Daily Report: header row5, data rows to total
wb=openpyxl.load_workbook(WB); ws=wb["Daily Report"]
last=5
for r in range(6,ws.max_row+1):
    if ws.cell(r,1).value not in (None,""): last=r
render_grid("Daily Report",5,last,ws.max_column,"preview_pp_daily.png","Daily Report \u2014 by zone (Sept 2026)",header_row=5)
