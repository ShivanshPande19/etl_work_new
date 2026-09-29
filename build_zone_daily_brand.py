#!/usr/bin/env python3
"""
Per-zone DAILY BRAND SALES report (September 2026), styled in the PetPooja
workbook's own colour scheme (Office blue + peach + light-yellow, Calibri).

Each zone = one sheet:
  rows  = each day of September
  cols  = DATE | DAY | <every brand's daily sales> | DAILY TOTAL
  bottom TOTAL row = brand-wise totals + the zone GRAND TOTAL.
"""

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from build_september_report import extract, inr, SRC, ZONES

OUT = "SEPTEMBER 2026 - ZONE DAILY BRAND REPORT.xlsx"
FONT = "Calibri"

# ---- PetPooja palette (extracted from the source workbook) ----
BLUE   = "2E5A87"   # header band / grand-total fill  (white text)
LBLUE  = "D9E1F2"   # daily-total column fill / banding
PEACH  = "FCE4D6"   # weekend row fill
YELLOW = "FFF2CC"   # total row fill
RED    = "C00000"   # accent text (grand total / weekend day)
GREY   = "666666"   # secondary text
INKTX  = "1F1F1F"   # primary text
LINE   = "BFBFBF"   # thin borders
DASHC  = "A6A6A6"   # dash for zero

INR = '[>=10000000]"\u20b9"##\\,##\\,##\\,##0;[>=100000]"\u20b9"##\\,##\\,##0;"\u20b9"#,##0'

hair = Side(style="thin", color=LINE)
box = Border(left=hair, right=hair, top=hair, bottom=hair)
med = Side(style="medium", color=BLUE)


def sty(cl, *, bold=False, size=11, color=INKTX, fill=None, align="center",
        fmt=None, border=box, italic=False, wrap=False, indent=0):
    cl.font = Font(name=FONT, bold=bold, size=size, color=color, italic=italic)
    cl.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap, indent=indent)
    if fill:
        cl.fill = PatternFill("solid", fgColor=fill)
    if fmt:
        cl.number_format = fmt
    if border is not None:
        cl.border = border
    return cl


def build_zone(wb, zone, brands, rows):
    ws = wb.create_sheet(zone.strip())
    ws.sheet_view.showGridLines = False
    ncol = 2 + len(brands) + 1
    dtot = ncol
    last = get_column_letter(ncol)

    # title (blue text, PetPooja style)
    ws.merge_cells(f"A1:{last}1")
    sty(ws.cell(1, 1), bold=True, size=16, color=BLUE, align="left", border=None,
        indent=1).value = f"{zone.strip()}  \u2014  DAILY BRAND SALES  \u2022  SEPTEMBER 2026"
    ws.row_dimensions[1].height = 26
    ws.merge_cells(f"A2:{last}2")
    sty(ws.cell(2, 1), size=10, color=GREY, italic=True, align="left", border=None,
        indent=1).value = "Day-wise net sales by brand  \u2022  Amounts in \u20b9 (INR)  \u2022  weekend rows shaded"
    ws.row_dimensions[2].height = 16
    ws.row_dimensions[3].height = 6

    # header
    hr = 4
    sty(ws.cell(hr, 1), bold=True, size=10.5, color="FFFFFF", fill=BLUE, wrap=True).value = "DATE"
    sty(ws.cell(hr, 2), bold=True, size=10.5, color="FFFFFF", fill=BLUE, wrap=True).value = "DAY"
    for j, b in enumerate(brands):
        sty(ws.cell(hr, 3 + j), bold=True, size=10, color="FFFFFF", fill=BLUE, wrap=True).value = b.strip().upper()
    sty(ws.cell(hr, dtot), bold=True, size=10.5, color="FFFFFF", fill=BLUE, wrap=True).value = "DAILY TOTAL"
    ws.row_dimensions[hr].height = 30

    brand_tot = {b: 0.0 for b in brands}
    grand = 0.0
    r = hr + 1
    for row in rows:
        wknd = row["day"] in ("Saturday", "Sunday")
        rf = PEACH if wknd else "FFFFFF"
        sty(ws.cell(r, 1), size=10.5, color=INKTX, fill=rf).value = row["date"].strftime("%d %b")
        sty(ws.cell(r, 2), size=10.5, bold=wknd, color=(RED if wknd else GREY), fill=rf).value = row["day"]
        dsum = 0.0
        for j, b in enumerate(brands):
            v = row["values"].get(b, 0.0)
            dsum += v
            brand_tot[b] += v
            c = ws.cell(r, 3 + j)
            if v:
                sty(c, size=10.5, color=INKTX, fill=rf, align="right", fmt=INR, indent=1).value = v
            else:
                sty(c, size=10.5, color=DASHC, fill=rf, align="center").value = "\u2013"
        sty(ws.cell(r, dtot), bold=True, size=10.5, color=INKTX, fill=LBLUE,
            align="right", fmt=INR, indent=1).value = dsum
        grand += dsum
        ws.row_dimensions[r].height = 18
        r += 1

    # TOTAL row
    sty(ws.cell(r, 1), bold=True, size=11, color=BLUE, fill=YELLOW,
        border=Border(left=hair, right=hair, top=med, bottom=hair)).value = "TOTAL"
    sty(ws.cell(r, 2), fill=YELLOW,
        border=Border(left=hair, right=hair, top=med, bottom=hair)).value = ""
    for j, b in enumerate(brands):
        sty(ws.cell(r, 3 + j), bold=True, size=10.5, color=INKTX, fill=YELLOW, align="right",
            fmt=INR, indent=1, border=Border(left=hair, right=hair, top=med, bottom=hair)).value = brand_tot[b]
    sty(ws.cell(r, dtot), bold=True, size=12, color="FFFFFF", fill=BLUE, align="right",
        fmt=INR, indent=1, border=Border(left=hair, right=hair, top=med, bottom=med)).value = grand
    ws.row_dimensions[r].height = 24

    # explicit GRAND TOTAL callout below
    r += 2
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=max(2, ncol - 2))
    sty(ws.cell(r, 1), bold=True, size=12, color="FFFFFF", fill=BLUE, align="left",
        indent=1).value = f"GRAND TOTAL  \u2014  {zone.strip()} (September 2026)"
    ws.merge_cells(start_row=r, start_column=ncol - 1, end_row=r, end_column=ncol)
    sty(ws.cell(r, ncol - 1), bold=True, size=13, color=RED, fill=YELLOW,
        align="right", fmt=INR, indent=1).value = grand
    ws.row_dimensions[r].height = 24

    # widths
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 12
    for j, b in enumerate(brands):
        w = max(max((len(x) for x in b.strip().split()), default=6) + 3,
                len(inr(brand_tot[b])) + 3, 13)
        ws.column_dimensions[get_column_letter(3 + j)].width = w
    ws.column_dimensions[last].width = max(len(inr(grand)) + 4, 15)
    ws.freeze_panes = ws.cell(hr + 1, 3)
    return grand


def main():
    src = load_workbook(SRC, data_only=True)
    out = Workbook()
    out.remove(out.active)
    total = 0
    for zone, meta in ZONES.items():
        brands, rows = extract(src[zone], meta["start"])
        g = build_zone(out, zone, brands, rows)
        total += g
        print(f"  {zone.strip():12s} brands={len(brands)} days={len(rows)} grand={inr(g)}")
    out.save(OUT)
    print("Saved:", OUT, "| all zones", inr(total))


if __name__ == "__main__":
    main()
