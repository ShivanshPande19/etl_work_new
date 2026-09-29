#!/usr/bin/env python3
"""Faithful preview of the all-months report: reads actual cell fills/fonts."""
import sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from openpyxl import load_workbook

WB = "MONTHLY SALES REPORT - ALL ZONES (PetPooja style).xlsx"


def argb_to_hex(color, default="FFFFFF"):
    try:
        if color is None or color.rgb is None or not isinstance(color.rgb, str):
            return "#" + default
        rgb = color.rgb
        if len(rgb) == 8:
            rgb = rgb[2:]
        return "#" + rgb
    except Exception:
        return "#" + default


def render(sheet, path, max_rows=None):
    wb = load_workbook(WB)
    ws = wb[sheet]
    ncol = ws.max_column
    nrow = ws.max_row if max_rows is None else min(ws.max_row, max_rows)

    # column widths (relative)
    widths = []
    for c in range(1, ncol + 1):
        cd = ws.column_dimensions[chr(64 + c)] if c <= 26 else None
        widths.append(cd.width if cd and cd.width else 12)
    total_w = sum(widths)

    fig_w = min(total_w * 0.13, 26)
    fig_h = min(nrow * 0.20 + 0.5, 60)
    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    ax.set_xlim(0, total_w)
    ax.set_ylim(0, nrow)
    ax.axis("off")
    ax.invert_yaxis()

    # precompute merged ranges to draw title bands across
    merged = list(ws.merged_cells.ranges)

    def merge_span(r, c):
        for mr in merged:
            if mr.min_row <= r <= mr.max_row and mr.min_col <= c <= mr.max_col:
                return mr
        return None

    drawn = set()
    for r in range(1, nrow + 1):
        for c in range(1, ncol + 1):
            if (r, c) in drawn:
                continue
            cell = ws.cell(r, c)
            mr = merge_span(r, c)
            if mr:
                x0 = sum(widths[:mr.min_col - 1]); x1 = sum(widths[:mr.max_col])
                for rr in range(mr.min_row, mr.max_row + 1):
                    for cc in range(mr.min_col, mr.max_col + 1):
                        drawn.add((rr, cc))
                anchor = ws.cell(mr.min_row, mr.min_col)
                cell = anchor
            else:
                x0 = sum(widths[:c - 1]); x1 = x0 + widths[c - 1]
            y0 = r - 1
            fill = "#FFFFFF"
            if cell.fill and cell.fill.fgColor and cell.fill.patternType == "solid":
                fill = argb_to_hex(cell.fill.fgColor)
            ax.add_patch(Rectangle((x0, y0), x1 - x0, 1, facecolor=fill,
                                   edgecolor="#BFBFBF", linewidth=0.4, zorder=1))
            val = cell.value
            if val is not None and val != "":
                if isinstance(val, (int, float)):
                    nf = cell.number_format or ""
                    if "20b9" in nf or "\u20b9" in nf or "#" in nf:
                        n = int(round(val)); s = str(abs(n))
                        if len(s) > 3:
                            last3 = s[-3:]; rest = s[:-3]; parts = []
                            while len(rest) > 2:
                                parts.insert(0, rest[-2:]); rest = rest[:-2]
                            if rest:
                                parts.insert(0, rest)
                            s = ",".join(parts) + "," + last3
                        txt = "\u20b9" + s
                    else:
                        txt = str(val)
                else:
                    txt = str(val)
                fcolor = argb_to_hex(cell.font.color, "1F1F1F") if cell.font and cell.font.color else "#1F1F1F"
                bold = bool(cell.font and cell.font.bold)
                sz = float(cell.font.size) if cell.font and cell.font.size else 11
                ha = cell.alignment.horizontal or "center"
                ha = {"left": "left", "right": "right", "center": "center"}.get(ha, "center")
                tx = x0 + 0.3 if ha == "left" else (x1 - 0.3 if ha == "right" else (x0 + x1) / 2)
                ax.text(tx, y0 + 0.5, txt, ha=ha, va="center",
                        color=fcolor, fontsize=min(sz * 0.62, 8.5),
                        fontweight="bold" if bold else "normal", zorder=2)
    plt.subplots_adjust(left=0.01, right=0.99, top=0.99, bottom=0.01)
    plt.savefig(path, dpi=130, bbox_inches="tight", facecolor="white")
    plt.close()
    print("wrote", path, f"({nrow} rows)")


if __name__ == "__main__":
    render("BENNETT", "preview_am_bennett.png")
    render("DISTRICT 9", "preview_am_district9.png")
    # first ~70 rows of a big zone to show 2 months stacked
    render("ALPHA 1", "preview_am_alpha1_top.png", max_rows=70)
