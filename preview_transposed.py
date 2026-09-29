#!/usr/bin/env python3
"""Preview of the transposed (brands x days) report."""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
from openpyxl import load_workbook

WB = "SEPTEMBER 2026 - ZONE SALES REPORT (BRANDS x DAYS).xlsx"
INK="#34515B"; INK2="#4A6D76"; RED="#4F7D6E"; REDT="#9BC3B6"
GREY="#8B969D"; GREYD="#5C6B72"; INKTX="#2F3E45"; LINE="#E7E2D9"
BAND="#F7F5F0"; FILL2="#EDF2EF"; WHITE="#FFFFFF"; DASH="#C4C0B6"


def inr(n):
    n=int(round(n)); s=str(abs(n)); sign="-" if n<0 else ""
    if len(s)>3:
        last3=s[-3:]; rest=s[:-3]; parts=[]
        while len(rest)>2: parts.insert(0,rest[-2:]); rest=rest[:-2]
        if rest: parts.insert(0,rest)
        g=",".join(parts)+","+last3
    else: g=s
    return f"{sign}\u20b9{g}"


def render(sheet, path):
    wb=load_workbook(WB); ws=wb[sheet]
    ncol=ws.max_column
    share_c=ncol; total_c=ncol-1
    header=[ws.cell(5,c).value for c in range(1,ncol+1)]
    data=[]; kinds=[]
    r=6
    while r<=ws.max_row:
        v1=ws.cell(r,1).value
        if v1 in (None,"") and all(ws.cell(r,c).value in (None,"") for c in range(1,ncol+1)):
            break
        rowvals=[ws.cell(r,c).value for c in range(1,ncol+1)]
        kind="total" if v1=="DAILY TOTAL" else "data"
        disp=[]
        for c,v in enumerate(rowvals):
            col=c+1
            if col==share_c and isinstance(v,(int,float)):
                disp.append(f"{v*100:.1f}%")
            elif isinstance(v,(int,float)):
                disp.append(inr(v))
            else:
                disp.append("" if v is None else str(v))
        data.append(disp); kinds.append(kind)
        if kind=="total": break
        r+=1

    # top brand = data row with max TOTAL
    top_row=-1; top_val=-1
    for i,(row,k) in enumerate(zip(data,kinds)):
        if k=="data":
            raw=ws.cell(6+i,total_c).value
            if isinstance(raw,(int,float)) and raw>top_val: top_val=raw; top_row=i

    # faithful relative column widths: brand wide, days uniform, total/share wider
    units=[2.6]+[1.0]*(ncol-3)+[1.7,1.15]
    tot_u=sum(units); colw=[u/tot_u for u in units]

    nrows=len(data)+1
    fig_w=min(0.66*ncol+4, 34)
    fig_h=0.46*nrows+2
    fig,ax=plt.subplots(figsize=(fig_w,fig_h)); ax.axis("off")
    ax.add_patch(Rectangle((0,1.0),1,0.11,transform=ax.transAxes,color=INK,clip_on=False,zorder=3))
    ax.add_patch(Rectangle((0,0.986),1,0.014,transform=ax.transAxes,color=RED,clip_on=False,zorder=3))
    zt=sheet.strip()
    ax.text(0.006,1.055,zt[0],transform=ax.transAxes,color=REDT,fontsize=16,fontweight="bold",va="center",zorder=4)
    ax.text(0.02,1.055,f"{zt[1:]}      SEPTEMBER 2026    \u2014    Brands x Days",transform=ax.transAxes,color="white",fontsize=16,fontweight="bold",va="center",zorder=4)

    tbl=ax.table(cellText=data, colLabels=[str(h) for h in header],
                 colWidths=colw, cellLoc="center", loc="center", bbox=[0,0,1,0.9])
    tbl.auto_set_font_size(False); tbl.set_fontsize(7.5)

    for (rr,cc),cell in tbl.get_celld().items():
        cell.set_edgecolor(LINE); cell.set_linewidth(0.5); txt=cell.get_text()
        if rr==0:
            hd=str(header[cc]) if cc<len(header) else ""
            wknd=("Sat" in hd) or ("Sun" in hd)
            cell.set_facecolor(RED if (cc+1)==total_c else INK2)
            txt.set_color(REDT if (wknd and (cc+1)!=total_c) else "white")
            txt.set_fontweight("bold"); txt.set_fontsize(7)
            cell.set_height(cell.get_height()*1.6)
            if cc==0: txt.set_ha("left")
            continue
        kind=kinds[rr-1]
        if kind=="total":
            cell.set_facecolor(RED if (cc+1)==total_c else INK)
            txt.set_color("white"); txt.set_fontweight("bold")
            if cc==0: txt.set_ha("left")
            elif cc>=1: txt.set_ha("right")
        else:
            cell.set_facecolor(BAND if (rr-1)%2 else WHITE)
            if cc==0:
                txt.set_color(INKTX); txt.set_fontweight("bold"); txt.set_ha("left")
            elif (cc+1)==total_c:
                cell.set_facecolor(FILL2); txt.set_color(INKTX); txt.set_fontweight("bold"); txt.set_ha("right")
            elif (cc+1)==share_c:
                cell.set_facecolor(FILL2); txt.set_ha("right")
                if (rr-1)==top_row: txt.set_color(RED); txt.set_fontweight("bold")
                else: txt.set_color(GREYD)
            elif data[rr-1][cc]=="\u2013":
                txt.set_color(DASH)
            else:
                txt.set_color(INKTX); txt.set_ha("right")
    plt.savefig(path,dpi=130,bbox_inches="tight",facecolor="white"); print("wrote",path)


render("CENTRAL 50","preview_T_central.png")
render("BENNETT","preview_T_bennett.png")
