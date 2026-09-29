#!/usr/bin/env python3
"""Render PNG previews of the generated report sheets (visual check)."""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

DARK="#0F2D3F"; BLUE="#16537E"; LTBLUE="#D6E4F0"; BAND="#F2F7FB"; WEEK="#FCEEEC"

def fmt(v):
    if v is None or v=="": return ""
    if isinstance(v,(int,float)): return f"\u20b9{v:,.0f}"
    return str(v)

def render(sheet, path, title):
    wb=load_workbook("SEPTEMBER 2026 - ZONE SALES REPORT.xlsx")
    ws=wb[sheet]
    # collect header (row4) + data rows until 'SHARE %' row
    ncol=ws.max_column
    hdr=[ws.cell(4,c).value for c in range(1,ncol+1)]
    body=[]; fills=[]
    r=5
    while r<=ws.max_row:
        first=ws.cell(r,1).value
        if first in ("SHARE %",): break
        if first is None and all(ws.cell(r,c).value in (None,"") for c in range(1,ncol+1)):
            break
        rowvals=[ws.cell(r,c).value for c in range(1,ncol+1)]
        body.append([fmt(v) for v in rowvals])
        day=ws.cell(r,2).value
        if first=="TOTAL": fills.append(LTBLUE)
        elif day in ("Saturday","Sunday"): fills.append(WEEK)
        elif (r-5)%2: fills.append(BAND)
        else: fills.append("white")
        if first=="TOTAL": break
        r+=1

    nrows=len(body)+1
    fig_h=0.34*nrows+1.2
    fig,ax=plt.subplots(figsize=(1.35*ncol+1, fig_h))
    ax.axis("off")
    ax.set_title(title, fontsize=15, fontweight="bold", color=DARK, pad=16)
    tbl=ax.table(cellText=body, colLabels=[str(h) for h in hdr],
                 cellLoc="center", loc="center")
    tbl.auto_set_font_size(False); tbl.set_fontsize(8.5)
    tbl.scale(1,1.3)
    for (rr,cc),cell in tbl.get_celld().items():
        cell.set_edgecolor("#D9E2EC")
        if rr==0:
            cell.set_facecolor(BLUE); cell.set_text_props(color="white",fontweight="bold")
        else:
            f=fills[rr-1]
            cell.set_facecolor(f)
            if body[rr-1][0]=="TOTAL":
                cell.set_text_props(fontweight="bold",color=DARK)
            if cc==ncol-1:
                cell.set_text_props(fontweight="bold")
    plt.tight_layout()
    plt.savefig(path,dpi=150,bbox_inches="tight")
    print("wrote",path)

render("ALPHA 1","preview_alpha1.png","ALPHA 1 — September 2026 Sales Report")
render("BENNETT","preview_bennett.png","BENNETT — September 2026 Sales Report")
