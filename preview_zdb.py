#!/usr/bin/env python3
"""Preview of the per-zone daily brand report (PetPooja colours)."""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

WB="SEPTEMBER 2026 - ZONE DAILY BRAND REPORT.xlsx"
BLUE="#2E5A87"; LBLUE="#D9E1F2"; PEACH="#FCE4D6"; YELLOW="#FFF2CC"
RED="#C00000"; GREY="#666666"; INKTX="#1F1F1F"; LINE="#BFBFBF"; DASHC="#A6A6A6"

def ind(n):
    n=int(round(n)); s=str(abs(n))
    if len(s)>3:
        last3=s[-3:]; rest=s[:-3]; parts=[]
        while len(rest)>2: parts.insert(0,rest[-2:]); rest=rest[:-2]
        if rest: parts.insert(0,rest)
        s=",".join(parts)+","+last3
    return "\u20b9"+s

def render(sheet, path):
    wb=load_workbook(WB); ws=wb[sheet]; ncol=ws.max_column
    header=[ws.cell(4,c).value for c in range(1,ncol+1)]
    data=[]; kinds=[]; days=[]
    r=5
    while r<=ws.max_row:
        v1=ws.cell(r,1).value
        if v1 in (None,""): break
        vals=[]
        for c in range(1,ncol+1):
            v=ws.cell(r,c).value
            vals.append(ind(v) if isinstance(v,(int,float)) else ("" if v is None else str(v)))
        data.append(vals); kinds.append("total" if v1=="TOTAL" else "day"); days.append(ws.cell(r,2).value)
        if v1=="TOTAL": break
        r+=1
    nrows=len(data)+1
    fig_w=min(1.35*ncol+1,22); fig_h=0.34*nrows+1.4
    fig,ax=plt.subplots(figsize=(fig_w,fig_h)); ax.axis("off")
    ax.set_title(f"{sheet.strip()}  —  Daily Brand Sales  •  September 2026", fontsize=14,
                 fontweight="bold", color=BLUE, loc="left", pad=14)
    tbl=ax.table(cellText=data, colLabels=[str(h) for h in header], cellLoc="center", loc="center", bbox=[0,0,1,0.95])
    tbl.auto_set_font_size(False); tbl.set_fontsize(8.5); tbl.scale(1,1.3)
    for (rr,cc),cell in tbl.get_celld().items():
        cell.set_edgecolor(LINE); cell.set_linewidth(0.6); txt=cell.get_text()
        if rr==0:
            cell.set_facecolor(BLUE); txt.set_color("white"); txt.set_fontweight("bold"); txt.set_fontsize(8)
            cell.set_height(cell.get_height()*1.4); continue
        k=kinds[rr-1]; day=days[rr-1]
        if k=="total":
            if cc==ncol-1: cell.set_facecolor(BLUE); txt.set_color("white")
            else: cell.set_facecolor(YELLOW); txt.set_color(BLUE if cc==0 else INKTX)
            txt.set_fontweight("bold")
        else:
            wknd=day in ("Saturday","Sunday")
            if cc==ncol-1: cell.set_facecolor(LBLUE); txt.set_fontweight("bold"); txt.set_color(INKTX)
            else: cell.set_facecolor(PEACH if wknd else "white")
            if cc==1: txt.set_color(RED if wknd else GREY); txt.set_fontweight("bold" if wknd else "normal")
            elif cc>=2 and data[rr-1][cc]=="\u2013": txt.set_color(DASHC)
            elif cc>=2: txt.set_color(INKTX)
        if cc>=2: txt.set_ha("right")
    plt.savefig(path,dpi=150,bbox_inches="tight",facecolor="white"); print("wrote",path)

render("CENTRAL 50","preview_zdb_central.png")
render("BENNETT","preview_zdb_bennett.png")
